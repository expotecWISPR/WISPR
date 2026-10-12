"""Genera docs/Informe_pruebas_seguridad_WISPR.docx (pruebas del 12 oct 2026).

Uso (requiere python-docx, fuera del venv del backend):
    python docs/generar_informe_seguridad.py docs/Informe_pruebas_seguridad_WISPR.docx
"""

import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = sys.argv[1]
AZUL = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x55, 0x55, 0x55)

doc = Document()
for section in doc.sections:
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = AZUL


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def p(text="", bold_prefix=None, style=None, size=None, italic=False, color=None, align=None, space_after=6):
    para = doc.add_paragraph(style=style)
    if bold_prefix:
        run = para.add_run(bold_prefix)
        run.bold = True
    run = para.add_run(text)
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if align:
        para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    return para


def bullets(items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            para.add_run(item[0]).bold = True
            para.add_run(item[1])
        else:
            para.add_run(item)
        para.paragraph_format.space_after = Pt(2)


def code(text):
    """Bloque de comandos/salida en fuente monoespaciada con fondo gris."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell, "F2F2F2")
    cell.paragraphs[0].text = ""
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(headers, rows, widths=None, header_fill="1F3A5F"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9.5)
            if i == len(row) - 1 and str(value).startswith(("Corregido", "Aprobado", "Pendiente", "Parche")):
                run.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def subsection(title):
    doc.add_heading(title, level=3)


# ---------------------------------------------------------------- Portada
for _ in range(5):
    doc.add_paragraph()
p("WISPR", size=34, color=AZUL, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0).runs[0].bold = True
p("Wireless Intelligence for Security, Practice and Response", size=12, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
doc.add_paragraph()
p("Informe de pruebas de seguridad", size=22, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2).runs[0].bold = True
p("Backend (FastAPI + Supabase) y prototipo de frontend", size=13, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(6):
    doc.add_paragraph()
for line in ["Colegio Técnico Profesional Don Bosco", "Especialidad de Ciberseguridad — Undécimo año", "Proyecto EXPOTEC 2026",
             "Fecha de ejecución de las pruebas: 12 de octubre de 2026"]:
    p(line, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
doc.add_page_break()

# ---------------------------------------------------------------- 1. Introducción
doc.add_heading("1. Introducción y alcance", level=1)
p("Este informe documenta las pruebas de seguridad solicitadas para WISPR, una plataforma educativa de tipo CTF "
  "(Capture The Flag) para practicar seguridad inalámbrica. Para cada prueba se indica el objetivo, el procedimiento "
  "seguido, el resultado obtenido, su análisis, la corrección aplicada cuando se encontró un problema y el resultado "
  "al repetir la prueba después de corregir.")
p("Componentes evaluados:", bold_prefix=None)
bullets([
    ("Backend: ", "API en Python (FastAPI) conectada a la base de datos PostgreSQL real en Supabase. Implementa el registro "
                  "e inicio de sesión con verificación por correo (MFA), los retos, el envío de flags, el ranking y el progreso."),
    ("Prototipo de frontend (wispr_terminal_professional): ", "interfaz web y servidor Node.js que todavía aloja la "
                  "terminal Kali en Docker. Se revisó su código y se prepararon correcciones verificadas."),
])
p("Entorno de prueba:", bold_prefix=None)
bullets([
    "Equipo Linux con el backend ejecutándose con uvicorn en 127.0.0.1:8000 (un solo proceso), contra la base de datos real.",
    "Las pruebas usaron cuentas creadas solo para este fin, que se eliminaron al terminar (la base quedó con 0 usuarios).",
    "Para no enviar correos a direcciones inventadas, el servidor de correo se desactivó durante las pruebas.",
    "Los comandos se muestran con curl de Linux. En Windows son idénticos usando curl.exe.",
])
p("Limitación: ", bold_prefix=None).runs[0].bold = True
doc.paragraphs[-1].add_run("el equipo de prueba no tiene Docker ni nmap. Por eso las pruebas de Contenedor, de Puertos con "
                           "nmap -sV y de Restauración completa quedan pendientes de ejecutar en la computadora donde se monta "
                           "WISPR completo. Para ellas se dejan los comandos exactos y el resultado esperado según el código.")

# ---------------------------------------------------------------- 2. Resumen
doc.add_heading("2. Resumen de resultados", level=1)
table(
    ["#", "Prueba", "Resultado inicial", "Resultado tras corregir", "Estado"],
    [
        ["1", "Permisos", "Rutas privadas protegidas; no se puede actuar como otro jugador", "Sin cambios en el backend; corregida la ruta de diagnóstico del prototipo", "Aprobado"],
        ["2", "Datos incorrectos", "2 errores 500 y nombre con HTML aceptado (XSS)", "Todo se rechaza con HTTP 400 y mensaje claro; XSS corregido", "Corregido"],
        ["3", "Fuerza bruta", "Bloqueo tras 5 fallos; sin límite por IP", "Bloqueo + límite por IP/correo + tiempo constante", "Corregido"],
        ["4", "Contenedor", "No ejecutable sin Docker", "Configuración revisada en el código", "Pendiente"],
        ["5", "Puertos", "Solo el puerto 8000 en 127.0.0.1", "Confirmado: no accesible desde la red", "Pendiente (nmap)"],
        ["6", "Dependencias", "36 vulnerabilidades (Python) + 1 alta (npm)", "0 vulnerabilidades en Python y 0 en npm", "Corregido"],
        ["7", "Restauración", "Falta el documento de la sección 12", "Medida la parte de base de datos", "Pendiente"],
    ],
    widths=[0.8, 2.6, 4.6, 5.2, 2.6],
)
p("Además se agregaron 23 pruebas automáticas nuevas: el backend pasó de 97 a 120 pruebas, y todas pasan.",
  italic=True, color=GRIS)

# ---------------------------------------------------------------- 3. Pruebas
doc.add_heading("3. Pruebas realizadas", level=1)

# --- 3.1 Permisos
doc.add_heading("3.1 Permisos", level=2)
subsection("Objetivo")
p("Verificar que un jugador con sesión iniciada no puede acceder a funciones o datos que no le corresponden, "
  "y que sin sesión no se puede acceder a ninguna ruta privada.")
subsection("Procedimiento")
code("""curl -i -H "Authorization: Bearer <token de jugador>" http://127.0.0.1:8000/api/kali/diagnose
curl -i http://127.0.0.1:8000/api/ranking                              # sin sesión
curl -i -H "Authorization: Bearer 0000...0000" http://127.0.0.1:8000/api/players/me   # token inventado
curl -H "Authorization: Bearer <token del jugador 7>" "http://127.0.0.1:8000/api/players/me?playerId=1\"""")
subsection("Resultado")
code("""/api/kali/diagnose con sesión     -> HTTP/1.1 404 Not Found
/api/ranking sin sesión           -> HTTP/1.1 401 Unauthorized
   {"ok":false,"auth":"required","error":"Tu sesión expiró o no es válida. Inicia sesión de nuevo."}
token inventado                   -> HTTP/1.1 401 Unauthorized
playerId=1 con la sesión del 7    -> responde con los datos del jugador 7 (@pruebas), no del 1""")
subsection("Análisis")
bullets([
    "Toda ruta privada exige un token de sesión válido. El servidor guarda solo el hash SHA-256 del token, y la sesión se puede revocar.",
    "El servidor ignora el identificador de jugador que manda el navegador y usa siempre el de la sesión. "
    "Cambiar un número en la URL no permite ver ni modificar datos de otra persona.",
    "/api/kali/diagnose responde 404 porque la terminal Kali todavía no se migró a este backend: sigue en el servidor Node del prototipo.",
    "Revisando el código del prototipo se encontró que esa ruta mostraba a cualquier jugador las versiones exactas de Docker, "
    "Node.js y el sistema operativo del servidor (información útil para un atacante), y que se podía pedir sin límite "
    "(cada pedido ejecuta unos 8 comandos de Docker).",
])
subsection("Corrección")
p("Se corrigió el servidor del prototipo (docs/parches/frontend-server.js.patch, aplicado): las versiones solo se "
  "muestran a administradores, el estudiante ve únicamente el estado de su propio contenedor, y se permite un diagnóstico "
  "cada 10 segundos por jugador. Resultado con una sesión real de estudiante, después de aplicar el cambio:")
code("""$ curl -H "Authorization: Bearer <token de estudiante>" http://127.0.0.1:8799/api/kali/diagnose
WISPR diagnóstico 2026-10-12T06:09:00.629Z
build kali=idle target=idle
FAIL red wispr-net-53fe4ed383a3ec            (FAIL: el equipo de prueba no tiene Docker)
FAIL contenedor wispr-kali-53fe4ed383a3ec
...                                          -> ya no aparecen versiones de Docker, Node ni del sistema
$ (segundo pedido inmediato)
{"ok":false,"error":"Espera unos segundos antes de volver a diagnosticar."}   [HTTP 429]""")
subsection("Estado")
p("Aprobado. La debilidad encontrada en el prototipo quedó corregida.")

# --- 3.2 Datos incorrectos
doc.add_heading("3.2 Datos incorrectos", level=2)
subsection("Objetivo")
p("Comprobar que el sistema rechaza datos inválidos o maliciosos sin fallar ni guardarlos: textos de 300 caracteres, "
  "código HTML/JavaScript (<script>alert(1)</script>), contraseña vacía y flag vacía.")
subsection("Procedimiento")
p("Se envió el registro (POST /api/players/register) con un campo inválido a la vez, y una flag vacía a "
  "POST /api/submissions con una sesión válida. Ejemplo:")
code("""curl -i -X POST http://127.0.0.1:8000/api/players/register -H "Content-Type: application/json" \\
  -d '{"name":"<script>alert(1)</script>","handle":"xss1","gender":"otro","password":"Clave12345","email":"xss1@example.com"}'""")
subsection("Resultado inicial")
table(
    ["Entrada", "Respuesta", "Evaluación"],
    [
        ["Nombre de 300 caracteres", "200, se recorta a 48", "Correcto"],
        ["Usuario de 300 caracteres", "200, se limpia a [a-z0-9._-] y se recorta a 32", "Correcto"],
        ["Correo de 300 caracteres", "200 al registrar, pero 500 al verificar el código", "Falla"],
        ["Contraseña de 300 caracteres", "500 Internal Server Error", "Falla"],
        ["Contraseña vacía", "400 «La contraseña debe tener al menos 4 caracteres.»", "Correcto (mínimo débil)"],
        ["<script> en el usuario", "Se limpia a @scriptalert1script", "Correcto"],
        ["<script> en el correo", "Aceptado", "Débil"],
        ["<script> / <img onerror> en el nombre", "Aceptado y guardado tal cual", "Falla (XSS)"],
        ["Flag vacía, espacios, null, WISPR{}", "«La flag debe tener el formato WISPR{...}.», no cuenta como intento", "Correcto"],
    ],
    widths=[4.6, 7.6, 3.6],
)
subsection("Análisis")
bullets([
    ("Contraseña larga (500): ", "el algoritmo bcrypt solo procesa 72 bytes. El código rechazaba las más largas lanzando "
                                  "un error, pero el registro no validaba el largo antes y el error salía como fallo interno (500). "
                                  "Lo mismo pasaba en el inicio de sesión contra una cuenta existente."),
    ("Correo largo (500): ", "la columna de la base de datos admite 255 caracteres. Un correo más largo pasaba el registro "
                             "y fallaba al crear la cuenta."),
    ("XSS almacenado: ", "el nombre se guardaba sin filtrar. Casi toda la interfaz lo muestra de forma segura, excepto el "
                         "gráfico «Puntos por estudiante» del profesor, que inserta el nombre con innerHTML sin escapar. "
                         "Se reprodujo en un navegador real: con un estudiante llamado <img src=x onerror=…>, el código del "
                         "atacante se ejecuta en la sesión del profesor, que podría perder su cuenta (el token está en localStorage)."),
])
subsection("Corrección")
bullets([
    "Contraseña: entre 8 y 64 caracteres y como máximo 72 bytes (el mínimo sube de 4 a 8, según NIST SP 800-63B). "
    "El inicio de sesión ya no falla con contraseñas largas.",
    "Correo: máximo 254 caracteres (estándar RFC 5321) y formato estricto. Ya no acepta <, >, comillas ni espacios.",
    "Nombre: lista de caracteres permitidos (letras de cualquier idioma, números, espacios, punto, guion y apóstrofo). "
    "Se rechaza cualquier etiqueta HTML.",
    "Frontend: el gráfico del profesor ahora escapa el nombre antes de insertarlo (docs/parches/frontend-app.js.patch, aplicado). "
    "Verificado en Chromium: antes el script del atacante se ejecutaba; ahora no, y el nombre aparece como texto. "
    "El formulario de registro también exige ahora entre 8 y 64 caracteres, igual que el servidor.",
])
subsection("Resultado tras la corrección")
code("""correo 300 caracteres   -> 400 {"ok":false,"error":"Escribe un correo válido."}
contraseña 300          -> 400 {"ok":false,"error":"La contraseña no puede tener más de 64 caracteres."}
<script> en el nombre   -> 400 {"ok":false,"error":"El nombre solo puede tener letras, números, espacios, puntos, guiones y apóstrofes."}
<img onerror> en nombre -> 400 (mismo mensaje)
<script> en el correo   -> 400 {"ok":false,"error":"Escribe un correo válido."}
contraseña vacía        -> 400 {"ok":false,"error":"La contraseña debe tener al menos 8 caracteres."}
flag vacía              -> {"ok":false,"error":"La flag debe tener el formato WISPR{...}."}
login con clave de 300  -> 401 {"ok":false,"error":"Usuario o contraseña incorrectos."}   (antes podía dar 500)""")
subsection("Estado")
p("Corregido en el backend (con pruebas automáticas que lo cubren) y en el frontend.")

# --- 3.3 Fuerza bruta
doc.add_heading("3.3 Fuerza bruta", level=2)
subsection("Objetivo")
p("Hacer 10 intentos de inicio de sesión fallidos seguidos y observar si el servidor bloquea o sigue respondiendo.")
subsection("Procedimiento")
code("""for i in $(seq 1 10); do
  curl -s -X POST http://127.0.0.1:8000/api/players/login -H "Content-Type: application/json" \\
       -d "{\\"handle\\":\\"pruebas\\",\\"password\\":\\"incorrecta$i\\"}"
done
# intento 11 con la contraseña CORRECTA""")
subsection("Resultado")
code("""intento  1..10: HTTP 401 {"ok":false,"error":"Usuario o contraseña incorrectos."}   (~0,6-1 s cada uno)
intento 11 (contraseña correcta): HTTP 401
   {"ok":false,"locked":true,"error":"Cuenta bloqueada temporalmente por intentos fallidos. Intenta de nuevo en ~15 min."}""")
subsection("Análisis")
bullets([
    "Tras 5 fallos en 15 minutos, la cuenta se bloquea 15 minutos. Mientras dura el bloqueo, ni la contraseña correcta permite entrar.",
    "El mensaje de error es el mismo exista o no la cuenta, así que no revela qué usuarios existen.",
    "Además, entrar exige el código de 6 dígitos enviado al correo (MFA): adivinar la contraseña no alcanza.",
    "Problemas encontrados: (1) no había límite por IP, así que un atacante podía probar unas pocas contraseñas en muchas "
    "cuentas distintas, o registrarse repetidamente con el correo de otra persona para llenarle la bandeja de códigos; "
    "(2) cuando el usuario no existía, el servidor respondía más rápido (se saltaba bcrypt), y midiendo el tiempo se podía "
    "averiguar qué cuentas existen.",
])
subsection("Corrección")
bullets([
    "Límite de peticiones en memoria: 60 intentos de inicio de sesión por minuto por IP, 30 envíos de código por minuto por IP, "
    "como máximo 3 registros al mismo correo cada 10 minutos, y 30 envíos de flags por minuto por jugador. Los límites por IP son "
    "generosos a propósito, porque en el colegio todos los estudiantes salen a internet con la misma IP pública.",
    "Tiempo constante: si el usuario no existe, se compara igual contra un hash de relleno.",
])
subsection("Resultado tras la corrección")
code("""tiempo promedio (5 intentos): cuenta existente 0,620 s  |  cuenta inexistente 0,612 s
ráfaga de 70 inicios de sesión desde una IP: 39 x HTTP 401 y 31 x HTTP 429
   {"ok":false,"error":"Demasiados intentos seguidos. Espera un minuto y vuelve a intentarlo."}
   (contando los 21 intentos previos, el corte ocurre exactamente en el intento 61)""")
subsection("Estado")
p("Corregido. El servidor bloquea la cuenta atacada, limita a quien hace muchos intentos y no revela qué cuentas existen.")

# --- 3.4 Contenedor
doc.add_heading("3.4 Contenedor", level=2)
subsection("Objetivo")
p("Verificar el aislamiento del contenedor Kali desde la terminal de WISPR: con qué usuario se ejecuta, si tiene acceso al "
  "socket de Docker (que permitiría controlar el host) y si puede alcanzar la red local.")
subsection("Procedimiento (pendiente de ejecutar)")
code("""whoami
ls -l /var/run/docker.sock
ping -c 2 <IP del router>""")
p("No se pudo ejecutar porque el equipo de prueba no tiene Docker. Debe correrse en la computadora donde se monta WISPR "
  "(Windows con Docker Desktop), abriendo un reto y usando la terminal integrada.")
subsection("Resultado esperado según el código")
p("Se revisó el comando docker run con el que el servidor crea el contenedor (server.js):")
table(
    ["Comando", "Resultado esperado", "Motivo en la configuración"],
    [
        ["whoami", "kali (no root)", "--user 1000:1000"],
        ["ls -l /var/run/docker.sock", "No such file or directory", "El socket no se monta. Solo /lab/challenges (solo lectura) y /workspace"],
        ["ping -c 2 <router>", "Falla (red inalcanzable u operación no permitida)", "Red Docker creada con --internal (sin salida) y --cap-drop ALL"],
    ],
    widths=[4.0, 5.0, 6.8],
)
p("Otras protecciones del contenedor: sistema de archivos de solo lectura (--read-only), --security-opt no-new-privileges, "
  "máximo 256 procesos, 2 GB de memoria y 2 CPU.")
subsection("Estado")
p("Pendiente de ejecución. Al correrla, comparar con la tabla anterior; cualquier diferencia es un hallazgo.")

# --- 3.5 Puertos
doc.add_heading("3.5 Puertos", level=2)
subsection("Objetivo")
p("Identificar qué servicios expone el equipo donde corre WISPR (nmap -sV localhost).")
subsection("Procedimiento")
p("nmap no está instalado en el equipo de prueba. Como aproximación se hizo un escaneo TCP de los 65 535 puertos de "
  "127.0.0.1 (conexión completa y lectura de la primera línea de respuesta), y luego se intentó conectar al backend "
  "desde la IP de red del equipo.")
subsection("Resultado")
code("""Puertos TCP abiertos en 127.0.0.1:
   5355/tcp  (sin banner)                     -> LLMNR, resolución de nombres del sistema operativo (no es WISPR)
   8000/tcp  HTTP/1.1 405 Method Not Allowed  -> backend WISPR (uvicorn)
  32164, 35871, 36645, 46079/tcp              -> puertos temporales del editor y otros procesos locales (no son WISPR)

Dirección de escucha del puerto 8000: 127.0.0.1 (LISTEN)
Desde la IP de red 192.168.1.49:8000 -> conexión rechazada""")
subsection("Análisis")
p("WISPR abre un único puerto (8000) y solo en la interfaz local, por lo que no es accesible desde otros equipos de la red. "
  "En el despliegue en AWS se mantiene así: el único servicio expuesto es nginx (HTTPS, 443), que reenvía a uvicorn de forma local. "
  "Además, uvicorn se configura para no anunciar su nombre en las respuestas (--no-server-header).")
subsection("Estado")
p("Resultado favorable, pero falta ejecutar nmap -sV localhost en el equipo donde se monta WISPR completo para el informe "
  "(incluye el servidor Node del prototipo y Docker).")

# --- 3.6 Dependencias
doc.add_heading("3.6 Dependencias", level=2)
subsection("Objetivo")
p("Detectar vulnerabilidades conocidas en las librerías de terceros (npm audit en el prototipo y pip-audit en el backend).")
subsection("Resultado inicial")
code("""pip-audit -r requirements.txt  ->  Found 69 known vulnerabilities in 7 packages  (36 únicas)
npm audit                      ->  nodemailer <=10.0.8  Severity: high  (3 avisos)  1 high severity vulnerability""")
table(
    ["Paquete", "Versión", "Vulnerab.", "Origen", "Acción"],
    [
        ["pillow", "10.4.0", "17", "together 1.x (indirecta)", "Eliminada: together 2.x no la usa"],
        ["starlette", "0.41.3", "7", "fastapi", "Actualizada a 1.6.0 (fastapi 0.141.1)"],
        ["python-multipart", "0.0.20", "6", "directa, sin uso", "Eliminada"],
        ["python-jose", "3.3.0", "3", "código JWT ya sin uso", "Eliminada junto con el código muerto"],
        ["ecdsa", "0.19.2", "1 (sin arreglo)", "python-jose", "Eliminada"],
        ["pytest", "8.3.4", "1", "directa (pruebas)", "Actualizada a 9.1.1"],
        ["scapy", "2.6.1", "1", "directa", "Actualizada a 2.7.0"],
        ["nodemailer (npm)", "10.0.2", "1 alta", "prototipo Node", "Parche: 10.0.11"],
    ],
    widths=[3.0, 1.8, 2.0, 3.8, 5.2],
)
subsection("Criterios de la corrección")
bullets([
    "Se prefirieron versiones publicadas hace al menos dos semanas, no las de esta semana, para dar tiempo a que aparezcan fallos.",
    "Antes de tocar el proyecto se probó todo en un entorno aparte: las 120 pruebas automáticas pasan con las versiones nuevas.",
    "Se verificó que la librería de Together AI 2.x mantiene la misma forma de llamada que usa el generador de retos.",
    "Se verificó con uv que los requisitos se instalan sin compilar nada en Windows y Linux, con Python 3.12, 3.13 y 3.14. "
    "Así se descubrió que pydantic 2.10.4 no se podía instalar en Python 3.14, y se actualizó a 2.13.5.",
])
subsection("Resultado tras la corrección")
code("""pip-audit sobre un entorno instalado desde cero con requirements.txt (58 paquetes):
    No known vulnerabilities found
npm audit en el prototipo, con nodemailer 10.0.11 instalado:
    found 0 vulnerabilities""")
subsection("Estado")
p("Corregido en el backend y en el prototipo. Recomendación: repetir ambas auditorías antes de cada entrega, porque "
  "cada semana se publican vulnerabilidades nuevas.")

# --- 3.7 Restauración
doc.add_heading("3.7 Restauración", level=2)
subsection("Objetivo")
p("Seguir los pasos de restauración (sección 12 del documento del proyecto) y medir el tiempo.")
subsection("Estado")
p("Pendiente: la sección 12 no está en el repositorio ni en los documentos del prototipo. Cuando se tenga, se ejecutan "
  "sus pasos y se mide el tiempo total. Como referencia, se midió la restauración de la base de datos del backend contra Supabase:")
code("""alembic -c migrations/alembic.ini upgrade head   # recrea las tablas            -> 5,4 s
python -m app.db.seed_challenges                 # recarga los 10 retos y archivos -> 6,8 s""")
p("Esto cubre solo la base de datos, no la restauración completa del sistema (servidor, Docker, frontend).", italic=True, color=GRIS)

# ---------------------------------------------------------------- 4. Parches prototipo
doc.add_heading("4. Correcciones aplicadas al prototipo", level=1)
p("El prototipo wispr_terminal_professional no está bajo control de versiones (Git). Por eso cada corrección se preparó "
  "y verificó primero sobre una copia, se guardó como archivo de parche en docs/parches/ y se aplicó el 12 de octubre "
  "después de respaldar los archivos originales. Después de aplicarlas se comprobó que el servidor arranca y responde, "
  "que la interfaz carga sin errores de JavaScript y que npm audit no reporta vulnerabilidades. Para aplicarlas en otra "
  "copia del prototipo:")
code("""git apply docs/parches/frontend-app.js.patch       # XSS del gráfico + contraseña de 8 a 64 caracteres
git apply docs/parches/frontend-server.js.patch    # /api/kali/diagnose sin versiones para estudiantes + límite
npm install nodemailer@10.0.11                     # vulnerabilidad alta de nodemailer""")

# ---------------------------------------------------------------- 5. Otras mejoras
doc.add_heading("5. Endurecimiento adicional del backend", level=1)
p("Además de lo que pedían las pruebas, se aplicaron estas medidas y se verificaron en modo producción:")
table(
    ["Medida", "Verificación"],
    [
        ["Documentación interactiva (/docs, /redoc, /openapi.json) oculta en producción", "HTTP 404 en las tres rutas"],
        ["Cabeceras de seguridad: nosniff, X-Frame-Options DENY, no-referrer, no-store, Content-Security-Policy", "Presentes en todas las respuestas de /api"],
        ["CORS limitado al dominio del frontend (configurable con CORS_ORIGINS)", "Origen ajeno: preflight 400; origen permitido: 200"],
        ["En producción, el código MFA nunca se devuelve en la respuesta ni se escribe en el registro", "Sin correo: HTTP 502, sin código"],
        ["Eliminación del código JWT sin uso", "Menos superficie de ataque y menos dependencias"],
    ],
    widths=[10.0, 5.8],
)

# ---------------------------------------------------------------- 6. Pendientes
doc.add_heading("6. Pendientes y recomendaciones", level=1)
bullets([
    "Ejecutar en el equipo con Docker las pruebas de Contenedor (3.4) y nmap -sV localhost (3.5), y adjuntar las salidas.",
    "Obtener la sección 12 y medir la restauración completa (3.7).",
    "Repetir pip-audit y npm audit antes de cada entrega.",
    "Al desplegar en AWS: HTTPS con certificado, HSTS en nginx y CORS_ORIGINS con el dominio real "
    "(pasos en la sección CONFIG AWS de docs/arquitectura_mvp.md).",
])

doc.save(OUT)
print("guardado:", OUT)
